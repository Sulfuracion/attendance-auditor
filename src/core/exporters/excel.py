"""Advanced OpenPyXL export engine with formula injection defense and conditional styling.

Generates professional, multi-sheet Excel reports with:
- Strict CSV/Excel formula injection escaping for text beginning with '=', '+', '-', '@'.
- Three dedicated worksheets: 'Resumen General', 'Detalle Fichajes', and 'Incidencias'.
- Corporate design: frozen header panes, auto-fitted column widths, and '0.00' numeric formatting.
- Conditional formatting rules (soft red for incidents, amber for excess, soft blue for night).
"""

import re
from io import BytesIO
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Formula injection detector regex
FORMULA_STARTER_REGEX = re.compile(r"^[\=\+\-\@\t\r]")

# Palette definitions
NAVY_HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

# Conditional Formatting Styles
INCIDENT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
INCIDENT_FONT = Font(name="Calibri", size=11, color="9C0006", bold=True)

EXCESS_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
EXCESS_FONT = Font(name="Calibri", size=11, color="9C6500")

NIGHT_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
NIGHT_FONT = Font(name="Calibri", size=11, color="1F497D")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def escape_excel_value(val: Any) -> Any:
    """Neutralize formula injection by prefixing suspicious string values with a single quote."""
    if isinstance(val, str) and FORMULA_STARTER_REGEX.match(val):
        return f"'{val}"
    return val


def auto_fit_columns(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Adjust worksheet column widths based on maximum cell content length."""
    for col in ws.columns:
        max_len = 0
        first_cell = col[0]
        col_letter = get_column_letter(first_cell.column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)


def format_header_row(ws: openpyxl.worksheet.worksheet.Worksheet, headers: list[str]) -> None:
    """Append and style the corporate header row."""
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = NAVY_HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def export_attendance_to_excel(
    reconciled_df: pd.DataFrame,
) -> BytesIO:
    """Generate a multi-sheet, secured, formatted Excel report in memory.

    Args:
        reconciled_df: Consolidated output DataFrame from reconcile_attendance_events.

    Returns:
        BytesIO object containing the complete .xlsx workbook.
    """
    wb = openpyxl.Workbook()

    # --------------------------------------------------------------------------
    # Sheet 1: Resumen General (Aggregated per employee)
    # --------------------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Resumen General"
    ws_summary.freeze_panes = "A2"

    summary_headers = [
        "Código Empleado",
        "Total Horas",
        "Horas Ordinarias",
        "Horas Nocturnas",
        "Horas Festivas",
        "Horas de Exceso",
        "Nº Incidencias",
    ]
    format_header_row(ws_summary, summary_headers)

    if not reconciled_df.empty:
        summary_df = (
            reconciled_df.groupby("employee_code", observed=False)
            .agg(
                total_hours=("total_hours", "sum"),
                ordinary_hours=("ordinary_hours", "sum"),
                night_hours=("night_hours", "sum"),
                holiday_hours=("holiday_hours", "sum"),
                excess_hours=("excess_hours", "sum"),
                incidents=("has_incident", "sum"),
            )
            .reset_index()
        )

        for _, row in summary_df.iterrows():
            ws_summary.append(
                [
                    escape_excel_value(row["employee_code"]),
                    float(row["total_hours"]),
                    float(row["ordinary_hours"]),
                    float(row["night_hours"]),
                    float(row["holiday_hours"]),
                    float(row["excess_hours"]),
                    int(row["incidents"]),
                ]
            )

        # Apply borders and numeric format
        for row_cells in ws_summary.iter_rows(
            min_row=2, max_row=ws_summary.max_row, min_col=1, max_col=7
        ):
            for idx, cell in enumerate(row_cells):
                cell.border = THIN_BORDER
                if idx in (1, 2, 3, 4, 5):
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif idx == 6:
                    cell.number_format = "0"
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")

    auto_fit_columns(ws_summary)

    # --------------------------------------------------------------------------
    # Sheet 2: Detalle Fichajes (Itemized day shifts and events)
    # --------------------------------------------------------------------------
    ws_detail = wb.create_sheet(title="Detalle Fichajes")
    ws_detail.freeze_panes = "A2"

    detail_headers = [
        "Fecha",
        "Código Empleado",
        "Entrada",
        "Salida",
        "Total Horas",
        "Horas Ordinarias",
        "Horas Nocturnas",
        "Horas Festivas",
        "Horas de Exceso",
        "Incidencia",
        "Tipo Incidencia",
        "Filas Origen",
    ]
    format_header_row(ws_detail, detail_headers)

    if not reconciled_df.empty:
        for _, row in reconciled_df.iterrows():
            date_val = str(row["date"]) if pd.notna(row["date"]) else "-"
            in_val = (
                row["punch_in"].strftime("%Y-%m-%d %H:%M:%S") if pd.notna(row["punch_in"]) else "-"
            )
            out_val = (
                row["punch_out"].strftime("%Y-%m-%d %H:%M:%S")
                if pd.notna(row["punch_out"])
                else "-"
            )
            has_inc_str = "SI" if row["has_incident"] else "NO"
            inc_type = str(row["incident_type"] or "-")
            src_rows = str(row["source_rows"] or "-")

            ws_detail.append(
                [
                    escape_excel_value(date_val),
                    escape_excel_value(row["employee_code"]),
                    in_val,
                    out_val,
                    float(row["total_hours"]),
                    float(row["ordinary_hours"]),
                    float(row["night_hours"]),
                    float(row["holiday_hours"]),
                    float(row["excess_hours"]),
                    has_inc_str,
                    escape_excel_value(inc_type),
                    escape_excel_value(src_rows),
                ]
            )

        max_row_det = ws_detail.max_row
        for row_cells in ws_detail.iter_rows(min_row=2, max_row=max_row_det, min_col=1, max_col=12):
            for idx, cell in enumerate(row_cells):
                cell.border = THIN_BORDER
                if idx in (4, 5, 6, 7, 8):
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif idx in (0, 1, 2, 3, 9, 10, 11):
                    cell.alignment = Alignment(horizontal="center")

        # Conditional Formatting Rules for Detalle Fichajes:
        if max_row_det >= 2:
            # 1. Soft red for Incidencia == "SI" (Column J: J2:J{max_row})
            rule_incident = CellIsRule(
                operator="equal",
                formula=['"SI"'],
                fill=INCIDENT_FILL,
                font=INCIDENT_FONT,
            )
            ws_detail.conditional_formatting.add(f"J2:J{max_row_det}", rule_incident)

            # 2. Soft amber for Horas de Exceso > 0 (Column I: I2:I{max_row})
            rule_excess = CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=EXCESS_FILL,
                font=EXCESS_FONT,
            )
            ws_detail.conditional_formatting.add(f"I2:I{max_row_det}", rule_excess)

            # 3. Soft blue for Horas Nocturnas > 0 (Column G: G2:G{max_row})
            rule_night = CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=NIGHT_FILL,
                font=NIGHT_FONT,
            )
            ws_detail.conditional_formatting.add(f"G2:G{max_row_det}", rule_night)

    auto_fit_columns(ws_detail)

    # --------------------------------------------------------------------------
    # Sheet 3: Incidencias (Exclusive anomalous records)
    # --------------------------------------------------------------------------
    ws_incidents = wb.create_sheet(title="Incidencias")
    ws_incidents.freeze_panes = "A2"

    incident_headers = [
        "Fecha",
        "Código Empleado",
        "Entrada",
        "Salida",
        "Tipo Incidencia",
        "Filas Origen",
        "Estado Revisión",
    ]
    format_header_row(ws_incidents, incident_headers)

    if not reconciled_df.empty:
        incidents_only = reconciled_df[reconciled_df["has_incident"]].copy()
        for _, row in incidents_only.iterrows():
            date_val = str(row["date"]) if pd.notna(row["date"]) else "-"
            in_val = (
                row["punch_in"].strftime("%Y-%m-%d %H:%M:%S") if pd.notna(row["punch_in"]) else "-"
            )
            out_val = (
                row["punch_out"].strftime("%Y-%m-%d %H:%M:%S")
                if pd.notna(row["punch_out"])
                else "-"
            )

            ws_incidents.append(
                [
                    escape_excel_value(date_val),
                    escape_excel_value(row["employee_code"]),
                    in_val,
                    out_val,
                    escape_excel_value(str(row["incident_type"] or "ANOMALIA")),
                    escape_excel_value(str(row["source_rows"] or "-")),
                    "PENDIENTE",
                ]
            )

        for row_cells in ws_incidents.iter_rows(
            min_row=2, max_row=ws_incidents.max_row, min_col=1, max_col=7
        ):
            for cell in row_cells:
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws_incidents)

    # Save to BytesIO stream
    output_stream = BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream


def generate_biometric_template() -> BytesIO:
    """Generate a clean, styled Excel (.xlsx) template for bulk biometric punch ingestion.

    Includes canonical column headers and guidance comments/examples for supervisors/admins.
    Columns:
    - codigo_empleado (str, e.g. EMP001)
    - fecha_hora (str, e.g. 2026-10-04 08:30:00)
    - tipo_evento (str, 'IN'/'ENTRADA' o 'OUT'/'SALIDA')

    Returns:
        BytesIO object containing the formatted .xlsx template.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet(title="Fichajes Biometricos")
    else:
        ws.title = "Fichajes Biometricos"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    headers = ["codigo_empleado", "fecha_hora", "tipo_evento"]
    format_header_row(ws, headers)

    # Sample rows showcasing valid formatting
    sample_data = [
        ["EMP001", "2026-10-01 08:00:00", "IN"],
        ["EMP001", "2026-10-01 16:30:00", "OUT"],
        ["EMP002", "2026-10-01 09:00:00", "ENTRADA"],
        ["EMP002", "2026-10-01 17:00:00", "SALIDA"],
        ["EMP003", "2026-10-01 22:00:00", "IN"],
        ["EMP003", "2026-10-02 06:00:00", "OUT"],
    ]

    for row in sample_data:
        ws.append([escape_excel_value(cell) for cell in row])

    # Style sample rows
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row_cells:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Calibri", size=11)

    auto_fit_columns(ws)

    output_stream = BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream


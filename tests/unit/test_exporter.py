"""Unit tests for the advanced OpenPyXL Excel export engine."""

from datetime import date
from io import BytesIO

import openpyxl
import pandas as pd

from src.core.exporters.excel import (
    escape_excel_value,
    export_attendance_to_excel,
    generate_biometric_template,
)


def test_escape_excel_value() -> None:
    """Validate that formula injection starters are properly escaped with a single quote."""
    assert escape_excel_value("=1+1") == "'=1+1"
    assert escape_excel_value("+SUM(A1:A5)") == "'+SUM(A1:A5)"
    assert escape_excel_value("-12345") == "'-12345"
    assert escape_excel_value("@cmd") == "'@cmd"
    assert escape_excel_value("\tTAB_ATTACK") == "'\tTAB_ATTACK"
    assert escape_excel_value("\rCARRIAGE_RETURN") == "'\rCARRIAGE_RETURN"

    # Non-formula strings and numbers should remain unmodified
    assert escape_excel_value("EMP001") == "EMP001"
    assert escape_excel_value("Normal text") == "Normal text"
    assert escape_excel_value(42.5) == 42.5
    assert escape_excel_value(None) is None


def test_export_attendance_to_excel_structure_and_formula_mitigation() -> None:
    """Validate sheet names, freeze panes, formula neutralization, and conditional formatting."""
    # Create sample reconciled DataFrame with edge cases and formula injection attempt
    reconciled_df = pd.DataFrame(
        {
            "employee_code": ["=SUM(A1:A10)", "EMP002", "EMP003"],
            "date": [date(2026, 1, 10), date(2026, 1, 11), date(2026, 1, 12)],
            "punch_in": [
                pd.Timestamp("2026-01-10 08:00:00"),
                pd.Timestamp("2026-01-11 22:00:00"),
                pd.NaT,
            ],
            "punch_out": [
                pd.Timestamp("2026-01-10 18:30:00"),
                pd.Timestamp("2026-01-12 06:00:00"),
                pd.Timestamp("2026-01-12 16:00:00"),
            ],
            "total_hours": [10.5, 8.0, 0.0],
            "ordinary_hours": [8.0, 8.0, 0.0],
            "night_hours": [0.0, 8.0, 0.0],
            "holiday_hours": [0.0, 0.0, 0.0],
            "excess_hours": [2.5, 0.0, 0.0],
            "has_incident": [False, False, True],
            "incident_type": [None, None, "ORPHAN_EXIT"],
            "source_rows": ["2,3", "4,5", "6"],
        }
    )

    excel_stream = export_attendance_to_excel(reconciled_df)
    assert isinstance(excel_stream, BytesIO)
    assert excel_stream.getbuffer().nbytes > 0

    # Load and inspect the generated workbook in memory
    wb = openpyxl.load_workbook(excel_stream, data_only=False)

    # 1. Verify expected worksheets
    assert wb.sheetnames == ["Resumen General", "Detalle Fichajes", "Incidencias"]

    # 2. Verify freeze panes on each sheet
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        assert ws.freeze_panes == "A2"

    # 3. Verify Formula Injection Escape in Detalle Fichajes
    ws_detail = wb["Detalle Fichajes"]
    # Row 2 is the first data row; Column B is Código Empleado
    emp_code_cell = ws_detail["B2"].value
    assert isinstance(emp_code_cell, str)
    assert emp_code_cell.startswith("'=")

    # 4. Verify Incidencias sheet contains only the incident row
    ws_incidents = wb["Incidencias"]
    # Max row should be 2 (1 header + 1 incident row)
    assert ws_incidents.max_row == 2
    assert ws_incidents["B2"].value == "EMP003"
    assert ws_incidents["E2"].value == "ORPHAN_EXIT"

    # 5. Verify conditional formatting rules on Detalle Fichajes
    cf_rules = ws_detail.conditional_formatting
    assert len(cf_rules) >= 3  # Rules for incident, excess hours, and night hours


def test_generate_biometric_template() -> None:
    """Validate biometric Excel template structure and compatibility with normalizer."""
    from src.core.parsers.normalizer import normalize_attendance_file

    template_stream = generate_biometric_template()
    assert isinstance(template_stream, BytesIO)
    assert template_stream.getbuffer().nbytes > 0

    wb = openpyxl.load_workbook(template_stream, data_only=True)
    assert "Fichajes Biometricos" in wb.sheetnames
    ws = wb["Fichajes Biometricos"]
    assert ws.freeze_panes == "A2"

    headers = [ws.cell(row=1, column=col).value for col in range(1, 4)]
    assert headers == ["codigo_empleado", "fecha_hora", "tipo_evento"]
    assert ws.max_row >= 6

    # Test roundtrip: Normalizer should successfully parse the generated template
    template_stream.seek(0)
    canonical_df = normalize_attendance_file(template_stream, file_format="xlsx")
    assert not canonical_df.empty
    assert list(canonical_df.columns) == ["employee_code", "timestamp", "event_type", "source_row"]
    assert "EMP001" in canonical_df["employee_code"].values
    assert set(canonical_df["event_type"].unique()).issubset({"IN", "OUT"})

